"""Tests for the core conversion pipeline."""

from __future__ import annotations

import base64
import io
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from html_to_pptx import convert, extract_measurements, render_pptx
from html_to_pptx.measurement_tree import walk_elements

FIXTURES_DIR = Path(__file__).parent / "fixtures"
SVG_RING_RGB = (0x38, 0xBD, 0xF8)
SVG_FADED_RGB = (0xF9, 0x73, 0x16)
SVG_RING_STROKE_CENTER_PX = 60  # stroke spans 18-102px of the 600px donut box


def _decode_png(src: str) -> Image.Image:
    return Image.open(io.BytesIO(base64.b64decode(src.split(",", 1)[1]))).convert("RGBA")


# ---------------------------------------------------------------------------
# End-to-end (these launch a browser — kept to a minimum)
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_convert_creates_valid_pptx(minimal_html: Path, tmp_pptx: Path):
    result = await convert(str(minimal_html), str(tmp_pptx))
    assert result == str(tmp_pptx)
    assert Path(result).stat().st_size > 0

    prs = Presentation(str(tmp_pptx))
    assert len(prs.slides) == 2


@pytest.mark.asyncio
async def test_file_not_found_raises():
    with pytest.raises(FileNotFoundError):
        await extract_measurements("/nonexistent/path.html")


@pytest.mark.asyncio
async def test_empty_html_raises(tmp_path: Path):
    empty = tmp_path / "empty.html"
    empty.write_text("<html><body></body></html>")
    with pytest.raises(ValueError, match="No slides found"):
        await extract_measurements(str(empty))


@pytest.mark.asyncio
async def test_inline_svg_raster_holds_only_the_svg():
    """HTML overlaid on an SVG is exported as its own text box, so baking it into
    the SVG picture too renders it twice (the "shadow text" bug)."""
    measurements = await extract_measurements(str(FIXTURES_DIR / "svg_overlay.html"))
    donut, faded = [e for e in walk_elements(measurements[0]["elements"]) if e.get("isImage")]

    ring = _decode_png(donut["src"])
    center = (ring.width // 2, ring.height // 2)
    assert ring.getpixel(center)[3] == 0, "overlaid label leaked into the SVG raster"
    assert ring.getpixel((0, 0))[3] == 0, "slide background leaked into the SVG raster"
    assert ring.getpixel((SVG_RING_STROKE_CENTER_PX, center[1])) == (*SVG_RING_RGB, 255)

    # render_pptx applies opacity to the picture, so the raster must not pre-blend it.
    assert faded["opacity"] == 0.5
    assert _decode_png(faded["src"]).getpixel((0, 0)) == (*SVG_FADED_RGB, 255)


# ---------------------------------------------------------------------------
# Measurement tests (reuse session-scoped fixture — no browser per test)
# ---------------------------------------------------------------------------


def test_measurements_returns_correct_slide_count(minimal_measurements: list[dict]):
    assert len(minimal_measurements) == 2


def test_measurements_have_required_keys(minimal_measurements: list[dict]):
    for slide in minimal_measurements:
        assert "index" in slide
        assert "elements" in slide
        assert "backgroundColor" in slide


def test_measurements_capture_text(minimal_measurements: list[dict]):
    def find_text(elements: list[dict], target: str) -> bool:
        for el in elements:
            if target in el.get("text", ""):
                return True
            if find_text(el.get("children", []), target):
                return True
        return False

    assert find_text(minimal_measurements[0]["elements"], "Title Slide")
    assert find_text(minimal_measurements[1]["elements"], "Content Slide")


def test_measurements_capture_background_color(minimal_measurements: list[dict]):
    assert minimal_measurements[0]["backgroundColor"] is not None


# ---------------------------------------------------------------------------
# Rendering tests (pure python-pptx, no browser)
# ---------------------------------------------------------------------------


def test_render_creates_correct_slide_count(minimal_measurements: list[dict]):
    prs = render_pptx(minimal_measurements)
    assert len(prs.slides) == 2


def test_render_sets_slide_dimensions(minimal_measurements: list[dict]):
    prs = render_pptx(minimal_measurements)
    assert abs(prs.slide_width - Inches(13.333)) < Inches(0.01)
    assert prs.slide_height == Inches(7.5)


def test_render_adds_shapes(minimal_measurements: list[dict]):
    prs = render_pptx(minimal_measurements)
    for slide in prs.slides:
        assert len(slide.shapes) > 0


def test_render_output_is_saveable(minimal_measurements: list[dict], tmp_pptx: Path):
    prs = render_pptx(minimal_measurements)
    prs.save(str(tmp_pptx))
    reloaded = Presentation(str(tmp_pptx))
    assert len(reloaded.slides) == 2
