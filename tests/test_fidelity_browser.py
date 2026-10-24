"""End-to-end fidelity checks against a real Chromium render of fixtures/fidelity.html."""

from __future__ import annotations

import base64
import io

import pytest
from PIL import Image
from pptx.enum.dml import MSO_FILL
from pptx.oxml.ns import qn

from html_to_pptx import extract_measurements, render_pptx
from html_to_pptx.measurement_tree import walk_elements

pytestmark = pytest.mark.asyncio(loop_scope="session")


def _find(slide: dict, predicate) -> dict:
    match = next((el for el in walk_elements(slide["elements"]) if predicate(el)), None)
    assert match is not None
    return match


def _with_text(slide: dict, text: str) -> dict:
    return _find(slide, lambda el: text in (el.get("text") or ""))


def _decode(data_uri: str) -> Image.Image:
    return Image.open(io.BytesIO(base64.b64decode(data_uri.split(",", 1)[1]))).convert("RGB")


async def test_every_slide_is_found(fidelity_measurements):
    assert len(fidelity_measurements) == 5


async def test_grid_slide_keeps_its_columns(fidelity_measurements):
    right = _with_text(fidelity_measurements[0], "Right column")
    assert right["x"] >= 960


async def test_reveal_animations_are_settled(fidelity_measurements):
    slide = fidelity_measurements[1]
    title = _with_text(slide, "Animated title")
    assert title["opacity"] == 1
    assert abs(title["y"] - 100) < 2  # final keyframe, not the translateY(40px) start
    assert _with_text(slide, "Faded paragraph")["opacity"] == 1


async def test_speaker_notes_are_extracted_not_rendered(fidelity_measurements):
    slide = fidelity_measurements[1]
    assert slide["notes"] == "Say hello to the board."
    assert not any("Say hello" in (el.get("text") or "") for el in walk_elements(slide["elements"]))


async def test_radial_slide_background_and_pseudo_glow_are_rasterized(fidelity_measurements):
    slide = fidelity_measurements[2]
    picture = _decode(slide["paintSrc"])
    width, height = picture.size
    red, green, _ = picture.getpixel((width // 2, height // 2))
    assert red > 200 and green < 60  # radial center
    assert picture.getpixel((width - 10, 10))[1] > 200  # ::before square, top right


async def test_paint_raster_keeps_text_native(fidelity_measurements):
    card = _with_text(fidelity_measurements[2], "Glow card text")
    assert card["paintSrc"]
    assert card["backgroundImage"] is None
    _, _, blue = _decode(card["paintSrc"]).getpixel((5, 5))
    assert blue > 60
    # The host's text is hidden during capture so it is not drawn twice.
    center = _decode(card["paintSrc"]).resize((1, 1)).getpixel((0, 0))
    assert center[0] < 40


async def test_canvas_and_webp_become_pictures(fidelity_measurements):
    slide = fidelity_measurements[3]
    canvas = _find(slide, lambda el: el.get("tag") == "canvas")
    assert _decode(canvas["src"]).getpixel((10, 10))[0] > 240
    webp = _find(slide, lambda el: el.get("tag") == "img")
    assert webp["src"].startswith(("data:image/png", "data:image/jpeg"))
    assert _decode(webp["src"]).getpixel((10, 10))[1] > 180


async def test_box_shadow_letter_spacing_and_row_rules_are_measured(fidelity_measurements):
    slide = fidelity_measurements[3]
    card = _find(slide, lambda el: (el.get("boxShadow") or "none") != "none")
    assert abs(card["width"] - 600) < 1
    assert _with_text(slide, "New feature")["letterSpacing"] == 4
    rule = _find(slide, lambda el: any(side["side"] == "bottom" for side in el.get("borderSides") or []))
    assert rule["borderSides"][0]["color"] == "rgb(255, 0, 255)"


async def test_quote_glyph_pseudo_is_painted_and_text_stays_native(fidelity_measurements):
    quote = _with_text(fidelity_measurements[3], "Design is how it works.")
    glyph = _decode(quote["paintSrc"])
    assert max(glyph.getpixel((x, y))[0] for x in range(40) for y in range(glyph.height)) > 200


async def test_fidelity_deck_renders_with_effects(fidelity_measurements, tmp_path):
    out = tmp_path / "fidelity.pptx"
    prs = render_pptx(fidelity_measurements)
    prs.save(out)
    xml = b"".join(slide.shapes._spTree.xml.encode() for slide in prs.slides)
    assert qn("a:outerShdw").split("}")[1].encode() in xml
    assert b'spc="200"' in xml
    assert b"NEW FEATURE" in xml
    assert prs.slides[1].notes_slide.notes_text_frame.text == "Say hello to the board."


async def test_preformatted_whitespace_is_kept_and_normal_whitespace_collapses(fidelity_measurements):
    slide = fidelity_measurements[4]
    assert _with_text(slide, "A  B")["text"] == "  A  B  "
    nested = _find(slide, lambda el: any(run["text"] == "  bold  " for run in el.get("inlineRuns") or []))
    assert "".join(run["text"] for run in nested["inlineRuns"]) == "  x   bold   y  "
    assert _with_text(slide, "collapsed")["text"] == "collapsed text\u00a0\u00a0kept"
    code = _find(slide, lambda el: el.get("tag") == "pre")
    assert [run["text"] for run in code["inlineRuns"]] == ["def f():", "\n", "    return 1"]
    byline = _find(slide, lambda el: any(run["text"] == "Dr. Amna" for run in el.get("inlineRuns") or []))
    assert [run["text"] for run in byline["inlineRuns"]] == ["Presented by ", "Dr. Amna", "\n", "Resident"]


async def test_visually_hidden_text_is_not_extracted(fidelity_measurements):
    texts = "".join(el.get("text") or "" for el in walk_elements(fidelity_measurements[4]["elements"]))
    assert "Skip to content" not in texts
    assert "Hidden label" not in texts


async def test_positive_thin_rules_survive_and_empty_boxes_do_not(fidelity_measurements):
    slide = fidelity_measurements[4]
    rules = [el for el in walk_elements(slide["elements"]) if el.get("backgroundColor") == "rgb(255, 0, 0)"]
    extents = sorted((el["width"], el["height"]) for el in rules)
    assert extents == [(0.5, 200), (1, 200), (200, 0.5), (200, 1), (200, 2)]

    prs = render_pptx([slide])
    red_shapes = [
        s for s in prs.slides[0].shapes
        if s.fill.type == MSO_FILL.SOLID and s.fill.fore_color.rgb == (255, 0, 0)
    ]
    assert len(red_shapes) == 5


async def test_preformatted_text_survives_render(fidelity_measurements):
    shapes = render_pptx([fidelity_measurements[4]]).slides[0].shapes
    texts = [s.text_frame.text for s in shapes if s.has_text_frame and s.text_frame.text]
    assert "  A  B  " in texts
    assert "  x   bold   y  " in texts
    assert "def f():\n    return 1" in texts


async def test_decks_without_slide_class_fall_back_to_sections(tmp_path):
    deck = tmp_path / "sections.html"
    deck.write_text(
        "<html><body style='margin:0'>"
        + "".join(
            f"<section style='width:1920px;height:1080px'><h1>Slide {n}</h1></section>" for n in range(3)
        )
        + "</body></html>"
    )
    measurements = await extract_measurements(str(deck))
    assert [_with_text(slide, "Slide")["text"] for slide in measurements] == ["Slide 0", "Slide 1", "Slide 2"]
