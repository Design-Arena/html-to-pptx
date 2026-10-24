"""Unit tests for OOXML effects, run typography, images and notes (no browser)."""

from __future__ import annotations

import base64
import io

from PIL import Image
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches

from html_to_pptx import render_pptx
from html_to_pptx.effects import parse_box_shadow

TAILWIND_SHADOW_LG = (
    "rgba(0, 0, 0, 0.1) 0px 10px 15px -3px, rgba(0, 0, 0, 0.1) 0px 4px 6px -4px"
)


def _slide(elements: list[dict], **slide_fields) -> list[dict]:
    return [{"index": 0, "backgroundColor": "rgb(255, 255, 255)", "elements": elements, **slide_fields}]


def _box(**fields) -> dict:
    return {
        "tag": "div", "x": 100, "y": 100, "width": 400, "height": 200, "text": "",
        "opacity": 1, "borderRadius": "0px", "children": [], **fields,
    }


def _webp_data_uri() -> str:
    buffer = io.BytesIO()
    Image.new("RGB", (8, 8), (200, 30, 30)).save(buffer, "WEBP")
    return "data:image/webp;base64," + base64.b64encode(buffer.getvalue()).decode()


def test_box_shadow_picks_the_furthest_reaching_outer_layer():
    shadow = parse_box_shadow(TAILWIND_SHADOW_LG + ", rgb(255, 0, 0) 0px 0px 40px 0px inset")
    assert (shadow.offset_y_px, shadow.blur_px, shadow.spread_px) == (10, 15, -3)
    assert abs(shadow.alpha - 0.1) < 1e-6


def test_box_shadow_none_and_inset_only():
    assert parse_box_shadow("none") is None
    assert parse_box_shadow("rgba(0, 0, 0, 0.2) 0px 2px 4px 0px inset") is None


def test_card_shadow_becomes_outer_shadow():
    prs = render_pptx(_slide([_box(backgroundColor="rgb(255, 255, 255)", boxShadow=TAILWIND_SHADOW_LG)]))
    shadow = prs.slides[0].shapes[0]._element.find(".//" + qn("a:outerShdw"))
    assert shadow is not None
    assert shadow.get("dir") == str(90 * 60000)  # straight down
    assert shadow.find(qn("a:srgbClr")).find(qn("a:alpha")).get("val") == "10000"


def test_shapes_without_css_shadow_do_not_inherit_the_theme_shadow():
    accent = _box(backgroundColor="rgb(241, 245, 249)", borderLeftColor="rgb(59, 130, 246)",
                  borderLeftWidth=4, borderLeftStyle="solid", borderRadius="0px 8px 8px 0px")
    shapes = list(render_pptx(_slide([accent])).slides[0].shapes)
    assert shapes
    for shape in shapes:
        assert shape._element.find(".//" + qn("a:outerShdw")) is None
        assert shape._element.find(f"{qn('p:style')}/{qn('a:effectRef')}").get("idx") == "0"


def test_letter_spacing_and_decoration_reach_the_run():
    el = _box(text="Badge", fontSize=20, letterSpacing=4, textDecoration="underline line-through",
              color="rgb(0, 0, 0)", fontFamily="Arial", height=600)
    prs = render_pptx(_slide([el]))
    r_pr = prs.slides[0].shapes[0]._element.find(".//" + qn("a:rPr"))
    assert r_pr.get("spc") == "200"  # 4px = 2pt = 200 hundredths
    assert r_pr.get("u") == "sng"
    assert r_pr.get("strike") == "sngStrike"


def test_side_borders_become_bars_at_true_width():
    divider = _box(borderSides=[
        {"side": "bottom", "width": 2, "color": "rgb(255, 0, 0)", "style": "solid"},
        {"side": "left", "width": 2, "color": "rgb(255, 0, 0)", "style": "solid"},
    ], borderLeftColor="rgb(255, 0, 0)", borderLeftWidth=2, borderLeftStyle="solid")
    shapes = list(render_pptx(_slide([divider])).slides[0].shapes)
    assert len(shapes) == 2
    bottom = max(shapes, key=lambda s: s.width)
    assert abs(bottom.top - Inches((100 + 200 - 2) * 7.5 / 1080)) < Inches(0.002)
    assert all(s.fill.fore_color.rgb == (255, 0, 0) for s in shapes)


def test_rounded_card_top_border_follows_corner_geometry():
    card = _box(borderRadius="16px", backgroundColor="rgb(255, 255, 255)", borderSides=[
        {"side": "top", "width": 4, "color": "rgb(0, 0, 255)", "style": "solid"},
    ])
    shapes = list(render_pptx(_slide([card])).slides[0].shapes)
    # Rounded accent underneath + inset card; no square bar poking past the corners.
    assert [s.auto_shape_type for s in shapes].count(shapes[0].auto_shape_type) == 2
    assert shapes[0].fill.fore_color.rgb == (0, 0, 255)


def test_webp_image_is_transcoded_instead_of_crashing():
    el = _box(tag="img", isImage=True, src=_webp_data_uri())
    pictures = list(render_pptx(_slide([el])).slides[0].shapes)
    assert len(pictures) == 1
    assert pictures[0].image.content_type == "image/png"


def test_missing_src_does_not_crash():
    render_pptx(_slide([_box(tag="img", isImage=True, src=None)]))


def test_speaker_notes_are_written(tmp_path):
    out = tmp_path / "notes.pptx"
    render_pptx(_slide([], notes="Mention the Q3 numbers.")).save(out)
    notes = Presentation(out).slides[0].notes_slide.notes_text_frame.text
    assert notes == "Mention the Q3 numbers."
