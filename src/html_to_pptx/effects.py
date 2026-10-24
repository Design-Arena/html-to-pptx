"""Box decoration and run typography that CSS expresses beyond fills.

* ``box-shadow`` → ``<a:outerShdw>``; everything else gets no shadow at all
  (python-pptx shapes otherwise inherit one from the theme)
* per-side borders (``border-bottom`` dividers, table rules) → thin bars
* ``letter-spacing`` / ``text-decoration`` → run ``spc`` / ``u`` / ``strike``
"""

from __future__ import annotations

import math
import re
from dataclasses import dataclass

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Inches

from .css_values import (
    RGB_RE,
    blend_over,
    css_color_to_rgb,
    rgb_match_to_color,
    split_top_level,
)
from .units import EMU_PER_CSS_PX, PIXELS_TO_INCHES_X, PIXELS_TO_INCHES_Y

_PX_RE = re.compile(r"(-?[\d.]+)px")
OOXML_PERCENT = 100_000
OOXML_DEGREE = 60_000
OOXML_HUNDREDTHS_PT = 100
OPAQUE_THRESHOLD = 0.99
# Thinner bars vanish in viewers; hairline CSS borders still need to show.
MIN_BORDER_BAR_INCHES = 0.007
# Elements of spPr that must follow <a:effectLst> (CT_ShapeProperties order).
_EFFECT_LST_SUCCESSORS = ("a:scene3d", "a:sp3d", "a:extLst")


@dataclass(frozen=True)
class BoxShadow:
    color: RGBColor
    alpha: float
    offset_x_px: float
    offset_y_px: float
    blur_px: float
    spread_px: float


def parse_box_shadow(css_value: str | None) -> BoxShadow | None:
    """Pick the most prominent outer layer of a computed ``box-shadow``.

    OOXML allows a single outer shadow per shape, so layered shadows (e.g.
    Tailwind's ``shadow-lg`` pair) collapse to the layer that reaches furthest.
    Inset layers have no faithful equivalent and are ignored.
    """
    if not css_value or css_value == "none":
        return None
    layers: list[BoxShadow] = []
    for layer in split_top_level(css_value):
        color_match = RGB_RE.search(layer)
        if "inset" in layer or not color_match:
            continue
        color, alpha = rgb_match_to_color(color_match)
        lengths = [float(v) for v in _PX_RE.findall(layer[color_match.end():])]
        if alpha < 0.01 or len(lengths) < 2:
            continue
        x, y, blur, spread = (lengths + [0.0, 0.0])[:4]
        layers.append(BoxShadow(color, alpha, x, y, blur, spread))
    if not layers:
        return None
    return max(layers, key=lambda s: abs(s.offset_x_px) + abs(s.offset_y_px) + s.blur_px + s.spread_px)


def apply_box_shadow(shape, shadow: BoxShadow, width_px: float, height_px: float) -> None:
    """Attach ``shadow`` to ``shape`` (autoshape or picture) as ``<a:outerShdw>``."""
    sp_pr = shape._element.spPr
    for old in sp_pr.findall(qn("a:effectLst")):
        sp_pr.remove(old)
    effect_lst = etree.Element(qn("a:effectLst"))
    outer = etree.SubElement(effect_lst, qn("a:outerShdw"))
    outer.set("blurRad", str(round(shadow.blur_px * EMU_PER_CSS_PX)))
    outer.set("dist", str(round(math.hypot(shadow.offset_x_px, shadow.offset_y_px) * EMU_PER_CSS_PX)))
    # OOXML measures dir clockwise from +x, which matches CSS's y-down offsets.
    direction_deg = math.degrees(math.atan2(shadow.offset_y_px, shadow.offset_x_px)) % 360
    outer.set("dir", str(round(direction_deg * OOXML_DEGREE)))
    if shadow.spread_px and width_px > 0 and height_px > 0:
        outer.set("sx", str(round((width_px + 2 * shadow.spread_px) / width_px * OOXML_PERCENT)))
        outer.set("sy", str(round((height_px + 2 * shadow.spread_px) / height_px * OOXML_PERCENT)))
    outer.set("algn", "ctr")
    outer.set("rotWithShape", "0")
    color = etree.SubElement(outer, qn("a:srgbClr"))
    color.set("val", str(shadow.color))
    if shadow.alpha < OPAQUE_THRESHOLD:
        etree.SubElement(color, qn("a:alpha")).set("val", str(round(shadow.alpha * OOXML_PERCENT)))
    sp_pr.insert_element_before(effect_lst, *_EFFECT_LST_SUCCESSORS)


def suppress_theme_shadows(slide) -> None:
    """Stop autoshapes from inheriting the theme's drop shadow.

    python-pptx autoshapes reference the theme effect style (``effectRef
    idx="2"``, a 35% outer shadow in the default template) that PowerPoint and
    LibreOffice both draw. Pointing the reference at ``idx="0"`` (no effect
    style, as PowerPoint's own default shapes do) removes it; an explicit
    ``<a:effectLst>`` from ``apply_box_shadow`` still takes precedence.
    LibreOffice ignores an empty ``<a:effectLst/>`` override, so that alone is
    not enough.
    """
    for shape in slide.shapes:
        effect_ref = shape._element.find(f"{qn('p:style')}/{qn('a:effectRef')}")
        if effect_ref is not None:
            effect_ref.set("idx", "0")


def apply_element_shadow(shape, el: dict) -> None:
    shadow = parse_box_shadow(el.get("boxShadow"))
    if shadow is not None:
        apply_box_shadow(shape, shadow, el.get("width", 0), el.get("height", 0))


def claim_top_border(el: dict) -> dict | None:
    """Return the element's lone top border as an accent and mark it drawn.

    On a rounded card a straight bar would poke square corners past the rounded
    top edge, so the caller draws it with the card's corner geometry instead.
    """
    for side in el.get("borderSides", []):
        if side["side"] == "top" and not side.get("_drawn"):
            side["_drawn"] = True
            return {"backgroundColor": side["color"], "height": side["width"]}
    return None


def mark_left_border_drawn(el: dict) -> None:
    for side in el.get("borderSides", []):
        if side["side"] == "left":
            side["_drawn"] = True


def render_side_borders(slide, el: dict, backdrop: RGBColor | None) -> None:
    """Draw non-uniform borders (dividers, table rules) as thin filled bars."""
    x, y = el.get("x", 0), el.get("y", 0)
    w, h = el.get("width", 0), el.get("height", 0)
    for side in el.get("borderSides", []):
        parsed = css_color_to_rgb(side["color"])
        if side.get("_drawn") or parsed is None:
            continue
        color, alpha = parsed
        if alpha < OPAQUE_THRESHOLD and backdrop is not None:
            color = blend_over(color, alpha, backdrop)
        t = side["width"]
        left_px, top_px, width_px, height_px = {
            "top": (x, y, w, t),
            "bottom": (x, y + h - t, w, t),
            "left": (x, y, t, h),
            "right": (x + w - t, y, t, h),
        }[side["side"]]
        bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(left_px * PIXELS_TO_INCHES_X),
            Inches(top_px * PIXELS_TO_INCHES_Y),
            Inches(max(width_px * PIXELS_TO_INCHES_X, MIN_BORDER_BAR_INCHES)),
            Inches(max(height_px * PIXELS_TO_INCHES_Y, MIN_BORDER_BAR_INCHES)),
        )
        bar.fill.solid()
        bar.fill.fore_color.rgb = color
        bar.line.fill.background()


def apply_run_typography(run, letter_spacing_pt: float, decoration: str | None) -> None:
    """Apply CSS letter-spacing and underline/line-through to a text run."""
    r_pr = run._r.get_or_add_rPr()
    if letter_spacing_pt:
        r_pr.set("spc", str(round(letter_spacing_pt * OOXML_HUNDREDTHS_PT)))
    if decoration and "underline" in decoration:
        r_pr.set("u", "sng")
    if decoration and "line-through" in decoration:
        r_pr.set("strike", "sngStrike")
